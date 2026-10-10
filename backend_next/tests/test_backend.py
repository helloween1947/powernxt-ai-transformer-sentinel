import io
import json
from datetime import datetime, timedelta, timezone
import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from powernxt.api import create_app
from powernxt.datasets import DatasetError, parse_dataset
from powernxt.engine import assess
from powernxt.schemas import Configuration, Measurements
from powernxt.simulator import generate
from powernxt.store import configuration

@pytest.fixture
def client(tmp_path):
    app = create_app(tmp_path / 'test.sqlite3', background=False)
    with TestClient(app) as client:
        assert client.post('/api/v1/assets', json={'asset_id': 'TX-01', 'name': 'Test transformer'}).status_code == 201
        yield client

def test_asset_configuration_validation_and_persistence(tmp_path):
    path = tmp_path / 'persist.sqlite3'
    with TestClient(create_app(path, background=False)) as c:
        assert c.post('/api/v1/assets', json={'asset_id': 'A', 'name': 'A', 'configuration': {'rated_kva': 0}}).status_code == 422
        assert c.post('/api/v1/assets', json={'asset_id': 'A', 'name': 'A'}).status_code == 201
        assert c.post('/api/v1/assets', json={'asset_id': 'A', 'name': 'A'}).status_code == 409
        assert c.put('/api/v1/assets/A/configuration', json={'rated_kva': 500}).json()['current_configuration']['version'] == 2
    with TestClient(create_app(path, background=False)) as c:
        assert c.get('/api/v1/assets/A').json()['current_configuration']['rated_kva'] == 500

def start(client, scenario='normal', asset='TX-01'):
    r = client.post('/api/v1/simulator/runs', json={'asset_id': asset, 'scenario': scenario, 'interval_seconds': .5})
    assert r.status_code == 201, r.text
    return r.json()

def test_stream_isolation_lifecycle_and_restart(client):
    a, b = start(client, 'overload'), start(client, 'normal')
    params = {'source': 'simulator', 'run_id': a['run_id']}
    latest = client.get('/api/v1/assets/TX-01/telemetry/latest', params=params).json()
    assert latest['source'] == 'simulator'
    analytics = client.get(f"/api/v1/telemetry/{latest['id']}/analytics").json()
    observations = analytics['result']['payload']['anomaly_observations']
    assert any(o['code'] == 'overload' and o['breached'] for o in observations)
    assert client.get('/api/v1/assets/TX-01/telemetry', params={'source': 'device'}).json()['items'] == []
    assert client.get('/api/v1/assets/TX-01/telemetry', params={'source': 'simulator'}).status_code == 422
    assert client.get('/api/v1/assets/TX-01/telemetry', params={'source': 'device', 'run_id': a['run_id']}).status_code == 422
    assert client.get('/api/v1/assets/TX-01/telemetry', params={'source': 'file_replay', 'run_id': a['run_id']}).status_code == 422
    assert client.patch(f"/api/v1/simulator/runs/{a['run_id']}", json={'status': 'stopped'}).json()['status'] == 'stopped'
    assert client.patch(f"/api/v1/simulator/runs/{a['run_id']}", json={'status': 'running', 'scenario': 'low_oil'}).json()['scenario'] == 'low_oil'
    run = client.app.state.store.run(a['run_id'])
    client.app.state.service.tick(run, datetime.now(timezone.utc) + timedelta(seconds=1))
    newer = client.get('/api/v1/assets/TX-01/telemetry/latest', params=params).json()
    assert newer['id'] > latest['id']
    assert newer['normalized_telemetry']['measurements']['oil_level_pct'] < 60
    assert client.get('/api/v1/assets/TX-01/telemetry/latest', params={'source': 'simulator', 'run_id': b['run_id']}).json()['id'] != newer['id']

@pytest.mark.parametrize('scenario,code', [('overload','overload'), ('overheating','overheating'), ('low_oil','low_oil'), ('phase_imbalance','phase_imbalance'), ('undervoltage','undervoltage')])
def test_synthetic_fault_scenarios_have_explainable_findings(client, scenario, code):
    run = start(client, scenario)
    faults = client.get('/api/v1/assets/TX-01/faults', params={'source':'simulator', 'run_id':run['run_id']}).json()
    found = [x for x in faults['items'] if x['code'] == code]
    assert found and all(x['breached'] and x['threshold'] is not None and x['reading_id'] for x in found)

def test_normal_zero_missing_and_threshold_boundaries():
    config = configuration(Configuration().model_dump(), 1, 'A')
    values = generate(config, 'normal', 42, 0, 2)
    assert not any(o['breached'] for o in assess(values, config)[1])
    values['oil_level_pct'] = 0
    values['oil_temperature_c'] = config['operational_limits']['max_top_oil_temp_c']
    observations = assess(values, config)[1]
    assert next(o for o in observations if o['code']=='low_oil')['breached'] is True
    assert next(o for o in observations if o['code']=='overheating')['breached'] is False
    missing = Measurements(oil_temperature_c=110).model_dump()
    observations = assess(missing, config)[1]
    assert next(o for o in observations if o['code']=='overload')['breached'] is None
    assert next(o for o in observations if o['code']=='overheating')['breached'] is True
    assert generate(config, 'normal', 42, 5, 2) == generate(config, 'normal', 42, 5, 2)

CSV = b'timestamp,oil_temperature_c,oil_level_pct\n2026-01-01T00:01:00Z,110,40\n2026-01-01T00:00:00Z,50,80\n'

def test_csv_upload_analysis_ordering_and_atomic_failure(client):
    result = client.post('/api/v1/assets/TX-01/datasets', files={'file': ('readings.csv', CSV, 'text/csv')})
    assert result.status_code == 201, result.text
    run = result.json()
    assert run['row_count'] == 2
    assert run['analysis']['counts'] == {'low_oil': 1, 'overheating': 1}
    params = {'source':'file_replay', 'run_id':run['run_id']}
    history = client.get('/api/v1/assets/TX-01/telemetry', params=params).json()['items']
    assert [r['measurement_time'] for r in history] == sorted([r['measurement_time'] for r in history], reverse=True)
    before = len(client.get('/api/v1/runs').json()['items'])
    bad = CSV + b'2026-01-01T00:02:00Z,NaN,50\n'
    assert client.post('/api/v1/assets/TX-01/datasets', files={'file': ('bad.csv', bad)}).status_code == 422
    assert len(client.get('/api/v1/runs').json()['items']) == before
    assert client.get('/api/v1/assets/TX-01/telemetry', params={**params,'limit':1,'offset':1}).json()['items'][0]['id'] == history[1]['id']
    filtered = client.get('/api/v1/assets/TX-01/telemetry', params={**params,'start':'2026-01-01T00:01:00Z'}).json()['items']
    assert len(filtered) == 1
    assert client.get('/api/v1/assets/TX-01/telemetry', params={**params,'start':'2026-01-01T00:02:00Z','end':'2026-01-01T00:01:00Z'}).status_code == 422

def test_xlsx_mapping_and_excel_datetime(client):
    book = Workbook(); sheet = book.active
    sheet.append(['Time','Oil','Level']); sheet.append([datetime(2026,1,1,12),111,45])
    stream = io.BytesIO(); book.save(stream)
    res = client.post('/api/v1/assets/TX-01/datasets', files={'file':('data.xlsx',stream.getvalue())}, data={'column_map':json.dumps({'timestamp':'Time','oil_temperature_c':'Oil','oil_level_pct':'Level'}),'timezone_name':'Asia/Kolkata'})
    assert res.status_code == 201, res.text
    run = res.json()
    assert run['analysis']['items'][0]['measurement_time'] == '2026-01-01T06:30:00+00:00'

@pytest.mark.parametrize('data,mapping', [
    (b'Oil\n110\n', {}),
    (b'timestamp,oil_temperature_c\n2026-01-01T00:00:00Z,110\n2026-01-01T00:00:00Z,120\n', {}),
    (b'timestamp,oil_level_pct\n2026-01-01T00:00:00Z,101\n', {}),
    (b'timestamp,oil_level_pct\n2026-01-01T00:00:00Z,inf\n', {}),
    (b'timestamp,oil_level_pct\n2099-01-01T00:00:00Z,50\n', {}),
    (CSV, {'timestamp':'not_here','oil_level_pct':'oil_level_pct'}),
    (b'timestamp,timestamp\na,b\n', {}),
])
def test_invalid_datasets_are_rejected(data,mapping):
    with pytest.raises(DatasetError):
        parse_dataset(data, 'bad.csv', mapping)

def test_device_ingestion_idempotency_and_config_snapshots(client):
    payload = {'message_id':'packet-1','timestamp':'2026-01-01T00:00:00Z','measurements':{'oil_temperature_c':110}}
    reading = client.post('/api/v1/assets/TX-01/telemetry',json=payload).json()
    assert reading['id']
    assert client.post('/api/v1/assets/TX-01/telemetry',json=payload).status_code == 409
    assert client.post('/api/v1/assets/TX-01/telemetry',json={**payload,'timestamp':'2099-01-01T00:00:00Z'}).status_code == 422
    client.put('/api/v1/assets/TX-01/configuration',json={'operational_limits':{'max_top_oil_temp_c':120}})
    result = client.get(f"/api/v1/telemetry/{reading['id']}/analytics").json()
    assert result['configuration_version'] == 1
    assert next(x for x in result['result']['payload']['anomaly_observations'] if x['code']=='overheating')['breached']


def test_asset_and_run_binding(client):
    run = start(client)
    client.post('/api/v1/assets',json={'asset_id':'B','name':'B'})
    assert client.get('/api/v1/assets/B/telemetry/latest',params={'source':'simulator','run_id':run['run_id']}).status_code == 422
    assert client.get('/api/v1/assets/missing').status_code == 404
    assert client.post('/api/v1/simulator/runs',json={'asset_id':'missing'}).status_code == 404
    assert client.get('/api/v1/assets/TX-01/telemetry',params={'limit':101}).status_code == 422


def test_dataset_summary_includes_faults_outside_first_page(client):
    rows = ['timestamp,oil_temperature_c']
    epoch = datetime(2026, 1, 1, tzinfo=timezone.utc)
    for i in range(205):
        rows.append(f'{(epoch+timedelta(minutes=i)).isoformat()},110')
    response = client.post('/api/v1/assets/TX-01/datasets', files={'file':('many.csv','\n'.join(rows).encode())})
    assert response.status_code == 201
    summary = response.json()['analysis']
    assert summary['readings_analyzed'] == 205
    assert summary['total_findings'] == 205
    assert summary['counts']['overheating'] == 205
    assert len(summary['items']) == 200 and summary['preview_truncated']
    assert summary['scope'] == 'entire_dataset'


def test_xlsx_preview_uses_active_sheet_and_rejects_bad_mapping(client):
    workbook = Workbook()
    workbook.active.append(['Wrong'])
    worksheet = workbook.create_sheet('Measurements'); worksheet.append(['Time','Oil'])
    worksheet.append([datetime(2026,1,1),110]); workbook.active = 1
    stream = io.BytesIO(); workbook.save(stream)
    file = ('multi.xlsx',stream.getvalue())
    response = client.post('/api/v1/assets/TX-01/datasets/preview', files={'file':file})
    assert response.json()['selected_sheet'] == 'Measurements'
    assert response.json()['columns'] == ['Time','Oil']
    assert client.post('/api/v1/assets/TX-01/datasets/preview',files={'file':file},data={'sheet':'Missing'}).status_code == 422
    assert client.post('/api/v1/assets/TX-01/datasets',files={'file':('broken.xlsx',b'bad file')}).status_code == 422
    assert client.post('/api/v1/assets/TX-01/datasets',files={'file':('x.csv',CSV)},data={'column_map':'no json'}).status_code == 422
    assert client.post('/api/v1/assets/TX-01/datasets',files={'file':('x.csv',CSV)},data={'column_map':'[]'}).status_code == 422


def test_dst_ambiguity_boolean_and_bad_timezone_are_rejected():
    with pytest.raises(DatasetError, match='Ambiguous'):
        parse_dataset(b'timestamp,oil_temperature_c\n2026-03-08T02:30:00,110\n','dst.csv',{},'America/New_York')
    with pytest.raises(DatasetError, match='timezone'):
        parse_dataset(CSV,'x.csv',{},'not-a-zone')
    with pytest.raises(ValueError, match='Boolean'):
        Measurements(oil_level_pct=True)


def test_failed_batch_rolls_back_readings_and_run(client):
    store = client.app.state.store
    rows = parse_dataset(CSV,'x.csv',{})
    rows[1]['message_id'] = rows[0]['message_id'] = 'duplicate'
    with pytest.raises(Exception):
        store.append('TX-01','file_replay','rollback',rows,{'run_id':'rollback','asset_id':'TX-01','source':'file_replay'})
    assert not store.history('TX-01','file_replay','rollback')
    with pytest.raises(KeyError):
        store.run('rollback')


def test_thermal_bootstrap_and_configuration_reset(client):
    payload = {'timestamp':'2026-01-01T00:00:00Z','message_id':'first','measurements':{'voltage_r_v':11000,'voltage_y_v':11000,'voltage_b_v':11000,'current_r_a':35,'current_y_a':35,'current_b_a':35,'ambient_temperature_c':28,'oil_temperature_c':50}}
    first = client.post('/api/v1/assets/TX-01/telemetry',json=payload).json()
    second = client.post('/api/v1/assets/TX-01/telemetry',json={**payload,'timestamp':'2026-01-01T00:01:00Z','message_id':'second'}).json()
    thermal = lambda reading: client.get(f"/api/v1/telemetry/{reading['id']}/analytics").json()['result']['payload']['thermal_assessment']
    assert thermal(first)['predicted_top_oil_temperature_c'] is None
    assert thermal(second)['predicted_top_oil_temperature_c'] > 50
    client.put('/api/v1/assets/TX-01/configuration',json={'rated_kva':2000})
    third = client.post('/api/v1/assets/TX-01/telemetry',json={**payload,'timestamp':'2026-01-01T00:02:00Z','message_id':'third'}).json()
    assert thermal(third)['predicted_top_oil_temperature_c'] is None


def test_upload_size_limits_and_unsupported_formats(client):
    assert client.post('/api/v1/assets/TX-01/datasets', files={'file':('data.csv', b'x'*(10*1024*1024+1))}).status_code == 422
    assert client.post('/api/v1/assets/TX-01/datasets', files={'file':('data.csv', b'x'*(12*1024*1024))}).status_code == 413
    assert client.post('/api/v1/assets/TX-01/datasets', files={'file':('data.xls', b'data')}).status_code == 422


def test_active_runs_and_readings_survive_restart(tmp_path):
    path = tmp_path/'restart.sqlite3'
    with TestClient(create_app(path,background=False)) as c:
        c.post('/api/v1/assets',json={'asset_id':'TX-01','name':'Restart test'})
        run = start(c)
        initial = c.get('/api/v1/assets/TX-01/telemetry/latest',params={'source':'simulator','run_id':run['run_id']}).json()
    with TestClient(create_app(path,background=False)) as c:
        assert c.get('/api/v1/runs').json()['items'][0]['status'] == 'running'
        stored = c.app.state.store.run(run['run_id'])
        c.app.state.service.tick(stored,datetime.fromisoformat(initial['measurement_time'])+timedelta(seconds=1))
        result = c.get('/api/v1/assets/TX-01/telemetry/latest',params={'source':'simulator','run_id':run['run_id']}).json()
        assert result['id'] > initial['id']
        assert c.get(f"/api/v1/telemetry/{initial['id']}/analytics").status_code == 200


def test_background_feed_advances_and_stops(tmp_path):
    import time
    with TestClient(create_app(tmp_path/'live.sqlite3')) as c:
        c.post('/api/v1/assets',json={'asset_id':'TX-01','name':'Live test'})
        run = start(c)
        params = {'source':'simulator','run_id':run['run_id']}
        initial = c.get('/api/v1/assets/TX-01/telemetry/latest',params=params).json()['id']
        time.sleep(.8)
        next_id = c.get('/api/v1/assets/TX-01/telemetry/latest',params=params).json()['id']
        assert next_id > initial
        c.patch(f"/api/v1/simulator/runs/{run['run_id']}",json={'status':'stopped'})
        time.sleep(.8)
        assert c.get('/api/v1/assets/TX-01/telemetry/latest',params=params).json()['id'] == next_id
        assert c.get('/health').json()['status'] == 'ok'
